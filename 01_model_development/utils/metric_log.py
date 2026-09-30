"""Write MIRROR metrics to the console and a document."""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


class MetricLogger:
    """Accumulate messages and metric summaries."""

    def __init__(self, destination):
        """Initialize the document and its destination."""
        self.destination = Path(destination)
        self.destination.parent.mkdir(parents=True, exist_ok=True)
        self.document = Document()
        normal = self.document.styles["Normal"]
        normal.font.name = "Consolas"
        normal.font.size = Pt(9)
        self.document.add_heading("MIRROR Metrics", level=0)
        self.document.add_paragraph(f"Created: {datetime.now():%Y-%m-%d %H:%M:%S}")

    def log(self, message):
        """Print a message and append it to the document."""
        print(message)
        self.document.add_paragraph(str(message))

    def log_results(self, results, thresholds):
        """Record grouped metrics and their thresholds."""
        metrics = ("n", "n_lab", "n_event", "acc", "auc", "sen", "spe", "c_index_hrd", "p_logrank", "hr")
        for modality, partitions in results.items():
            self.document.add_heading(f"{modality} | Threshold = {thresholds[modality]:.4f}", level=2)
            table = self.document.add_table(rows=1, cols=len(metrics) + 1, style="Light Shading Accent 1")
            for cell, label in zip(table.rows[0].cells, ("Set",) + metrics):
                cell.text = label
            for partition, values in partitions.items():
                cells = table.add_row().cells
                cells[0].text = partition
                for cell, metric in zip(cells[1:], metrics):
                    value = values[metric]
                    cell.text = "n/a" if value == -1 else f"{value:.4g}"
                self.log(f"{modality}/{partition}: n={values['n']}, AUC={values['auc']:.4f}")

    def save(self):
        """Save the accumulated document."""
        self.document.save(self.destination)
